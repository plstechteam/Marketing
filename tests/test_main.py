"""Tests for the pure helpers in main.py. Run: python tests/test_main.py"""
import os
import sys
from datetime import date, datetime

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
import main  # noqa: E402

P = main.PACIFIC


def test_month_windows_cover_the_range_without_gaps():
    start = P.localize(datetime(2026, 1, 1))
    end = P.localize(datetime(2026, 9, 23, 10, 30))
    w = main.month_windows(start, end)
    assert len(w) == 9
    assert w[0][0] == start and w[-1][1] == end
    for (_, a_end), (b_start, _) in zip(w, w[1:]):
        assert a_end == b_start
    # March crosses DST; the cut is still local midnight on the 1st.
    assert w[3][0] == P.localize(datetime(2026, 4, 1))


def test_month_windows_december_rolls_over():
    start = P.localize(datetime(2026, 12, 1))
    end = P.localize(datetime(2027, 1, 5))
    assert main.month_windows(start, end)[0][1] == P.localize(datetime(2027, 1, 1))


def test_label_value_maps_single_and_multi_and_keeps_unknown():
    labels = {"VDS": "Agency", "Scorpion": "SEO: Scorpion"}
    assert main.label_value("VDS", labels) == "Agency"
    assert main.label_value("VDS;Scorpion", labels) == "Agency;SEO: Scorpion"
    assert main.label_value("Gone", labels) == "Gone"
    assert main.label_value("", labels) is None
    assert main.label_value(None, labels) is None
    # People dropdowns: an archived owner has no option but is in the owner map.
    people = {"77": "Ann Lee"}
    assert main.label_value("77", people, {"99": "Old Owner"}) == "Ann Lee"
    assert main.label_value("99", people, {"99": "Old Owner"}) == "Old Owner"
    assert main.label_value("12", people, {"99": "Old Owner"}) == "12"


def test_datetime_goes_to_california():
    # PST in winter, PDT in summer.
    assert main.parse_hubspot_datetime("2026-01-01T03:00:00.000Z") == datetime(2025, 12, 31, 19, 0)
    assert main.parse_hubspot_datetime("2026-07-01T03:00:00.000Z") == datetime(2026, 6, 30, 20, 0)
    assert main.parse_hubspot_datetime(None) is None
    assert main.parse_hubspot_date("2026-02-03") == date(2026, 2, 3)


def test_stage_date_properties_falls_back_to_exited_then_skips():
    defs = {"hs_v2_date_entered_1": {}, "hs_v2_date_exited_1": {}, "hs_v2_date_exited_2": {}}
    cols, skipped = main.stage_date_properties(["1", "2", "3"], defs)
    assert cols == ["hs_v2_date_entered_1", "hs_v2_date_exited_2"]
    assert skipped == ["3"]


def test_close_date_for_picks_the_field_that_matches_how_it_closed():
    from datetime import date
    settled = {"date___settled": "2026-05-04", "date___dropped": "2026-01-01",
               "hs_v2_date_entered_current_stage": "2026-06-01T12:00:00Z"}
    assert main.close_date_for(settled, "Settled - Pre Lit", True) == (date(2026, 5, 4), "date___settled")
    closed_out = {"date___dropped": "2026-02-03"}
    assert main.close_date_for(closed_out, "Close Out", True) == (date(2026, 2, 3), "date___dropped")
    after_retained = {"date__intake_sign_up_close_out": "2026-03-01"}
    assert main.close_date_for(after_retained, "Retained - Client Dropped", True) == \
        (date(2026, 3, 1), "date__intake_sign_up_close_out")
    # Closed with the firm's field blank: never left undated.
    bare = {"hs_v2_date_entered_current_stage": "2026-06-01T12:00:00Z"}
    assert main.close_date_for(bare, "Settled - Lit", True) == \
        (date(2026, 6, 1), "hs_v2_date_entered_current_stage")
    # Settled while still in a Retained stage: Date - Settled wins.
    retained = {"date___settled": "2026-09-15", "hs_v2_date_entered_current_stage": "2026-08-01T12:00:00Z"}
    assert main.close_date_for(retained, "Retained - Pre Lit", True) == (date(2026, 9, 15), "date___settled")
    referred = {"date___referred_out": "2026-04-09"}
    assert main.close_date_for(referred, "Referred Out - Complete", True) == \
        (date(2026, 4, 9), "date___referred_out")
    assert "Referred Out - Complete" in main.CLOSED_STAGE_LABELS
    # Open: no close date even when a date field happens to be filled.
    assert main.close_date_for(settled, "Intake", False) == (None, None)


def test_closed_stages_include_close_out_and_fail_when_renamed():
    assert "Close Out" in main.CLOSED_STAGE_LABELS
    everything = [{"label": l} for l in main.CLOSED_STAGE_LABELS]
    main.check_closed_stages(everything)          # no exit
    try:
        main.check_closed_stages(everything[1:])
    except SystemExit:
        pass
    else:
        raise AssertionError("a missing closed stage must stop the run")


def test_parse_number():
    assert main.parse_number("5500") == 5500.0
    assert main.parse_number("12.5") == 12.5
    assert main.parse_number("") is None
    assert main.parse_number("n/a") is None


def test_ab1755_maps_stored_manufacturer_values():
    assert main.ab1755_for("General Motors LLC") == "Opt In"
    assert main.ab1755_for("Toyota Motor Sales, U.S.A., Inc. / Lexus") == "Opt Out"
    assert main.ab1755_for("Winnebago Industries, Inc.") == "Not on list"
    assert main.ab1755_for(None) == "Not on list"
    assert main.ab1755_for("") == "Not on list"
    assert not (main.AB1755_OPT_IN & main.AB1755_OPT_OUT)


def test_sheet_lookups_order_and_headers():
    props = ["hs_object_id", "dealname", "pipeline", "dealstage", "hubspot_owner_id",
             "createdate", "lead___source", "s__manufacturer", "ro_review__final_decision_",
             "hs_v2_date_entered_5411633", "some_new_property"]
    defs = {
        "dealname": {"label": "Deal Name", "type": "string"},
        "pipeline": {"label": "Pipeline", "type": "enumeration"},
        "dealstage": {"label": "Deal Stage", "type": "enumeration"},
        "createdate": {"label": "Create Date", "type": "datetime"},
        "lead___source": {"label": "Lead - Source", "type": "enumeration",
                          "options": [{"value": "VDS", "label": "Agency"}]},
        "s__manufacturer": {"label": "Manufacturer", "type": "enumeration",
                            "options": [{"value": "Polestar", "label": "POLESTAR AUTOMOTIVE USA, INC."}]},
        "ro_review__final_decision_": {"label": "RO Review (Final Decision)", "type": "enumeration",
                                       "options": [{"value": "Sign Up - Pre-Lit (OPT IN OEM)",
                                                    "label": "Pre-Lit (OPT IN OEM)"}]},
        "hs_v2_date_entered_5411633": {"label": 'Date entered "Intake (Lemon Law)"', "type": "datetime"},
        "some_new_property": {"label": "Some New Property", "type": "string"},
        "date___dropped": {"label": "Date - Closed Out", "type": "date"},
    }
    deals = [{"id": "1", "properties": {
        "hs_object_id": "1", "dealname": "Doe, Jane", "pipeline": "default", "dealstage": "5411635",
        "hubspot_owner_id": "77", "createdate": "2026-03-01T12:00:00Z", "lead___source": "VDS",
        "s__manufacturer": "Polestar", "ro_review__final_decision_": "Sign Up - Pre-Lit (OPT IN OEM)",
        "hs_v2_date_entered_5411633": None, "some_new_property": "x", "date___dropped": "2026-03-05"}}]
    stages = [{"id": "5411633", "label": "Intake", "displayOrder": 0},
              {"id": "5411635", "label": "Close Out", "displayOrder": 12}]
    df = main.build_deals_frame(deals, props, defs, {s["id"]: s["label"] for s in stages},
                                {"77": "Ann Lee"}, "Lemon Law")
    df = main.add_stage_attributes(df, stages)
    df = main.add_close_date(df, deals, stages)
    df["last_refresh"] = datetime(2026, 9, 23)
    sheet = main.to_sheet(df, defs, ["hs_v2_date_entered_5411633"])
    assert list(sheet.columns) == [
        "Record ID", "Deal Name", "Create Date",
        "Deal Stage", "Deal Stage ID", "Deal Stage Order", "Is Closed", "Is Settled",
        "Close Date", "Close Date Source",
        "Lead - Source",
        "Manufacturer", main.AB1755_HEADER, "RO Review (Final Decision)",
        "Deal Owner", "Deal Owner ID",
        'Date entered "Intake (Lemon Law)"',  # raw stamp, Stage history
        "Some New Property",                  # unplaced property: kept, before Audit
        "Pipeline", "Last Refresh"]
    row = sheet.iloc[0]
    assert row["Pipeline"] == "Lemon Law"
    assert row["Deal Stage"] == "Close Out" and row["Deal Stage ID"] == "5411635"
    assert row["Deal Stage Order"] == 12 and row["Is Closed"] == True  # noqa: E712
    assert row["Is Settled"] == False  # noqa: E712
    assert row["Close Date"] == date(2026, 3, 5)
    assert row["Close Date Source"] == "Date - Closed Out"   # label, not internal name
    assert row["Deal Owner"] == "Ann Lee" and row["Deal Owner ID"] == "77"
    assert row["Lead - Source"] == "Agency"
    assert row["Manufacturer"] == "POLESTAR AUTOMOTIVE USA, INC."
    assert row[main.AB1755_HEADER] == "Opt Out"
    assert row["Create Date"] == datetime(2026, 3, 1, 4, 0)   # 12:00 UTC = 4 AM PST


def test_every_base_property_has_a_place_in_the_layout():
    placed = {k for _, keys in main.COLUMN_GROUPS for k in keys}
    unplaced = [p for p in main.BASE_PROPERTIES if p not in placed]
    assert not unplaced, f"add to COLUMN_GROUPS: {unplaced}"
    assert set(main.FIXED_HEADERS) <= placed


def test_durations_are_calendar_days_and_blank_when_a_date_is_missing():
    import pandas as pd
    assert main.days_between(datetime(2026, 1, 1, 23, 0), date(2026, 1, 3)) == 2
    assert main.days_between(date(2026, 1, 3), datetime(2026, 1, 1, 1, 0)) == -2
    assert main.days_between(None, date(2026, 1, 3)) is None
    df = pd.DataFrame({
        "createdate": [datetime(2026, 1, 1, 9), datetime(2026, 2, 1, 9)],
        main.INTAKE_DATE: [datetime(2026, 1, 2, 9), None],
        "date___ro_review": [date(2026, 1, 11), None],
        "date___retained": [date(2026, 1, 15), None],
        main.READY_FOR_LEGAL: [datetime(2026, 1, 21, 15), None],
        "date___settled": [date(2026, 5, 1), None],
    })
    out = main.add_durations(df)
    first = {k: out[k].tolist()[0] for k, _, _, _ in main.DURATIONS}
    assert first == {
        "days_created_to_intake": 1,
        "days_created_to_ro_review": 10,
        "days_intake_to_signed": 13,
        "days_ro_review_to_signed": 4,
        "days_signed_to_ready_for_legal": 6,
        "days_ro_review_to_ready_for_legal": 10,
        "days_intake_to_ready_for_legal": 19,
        "days_created_to_ready_for_legal": 20,
        "days_created_to_settled": 120,
    }
    assert out["days_intake_to_signed"].isna().tolist() == [False, True]


def test_milestones_in_case_order():
    cols = ["createdate", main.INTAKE_DATE, "intake_date_source", main.INTAKE_LEGACY,
            "date___ro_review", "date___retained", main.READY_FOR_LEGAL,
            main.INTAKE_ENTERED, "hs_v2_date_entered_5411640"]
    order = main.column_order(cols, [main.INTAKE_ENTERED, "hs_v2_date_entered_5411640"])
    assert order.index(main.INTAKE_DATE) < order.index("date___ro_review") \
        < order.index("date___retained") < order.index(main.READY_FOR_LEGAL) \
        < order.index(main.INTAKE_ENTERED)          # the raw stamp is back in Stage history


def test_intake_date_prefers_the_stage_stamp_and_backfills_blanks():
    import pandas as pd
    df = pd.DataFrame({
        main.INTAKE_ENTERED: [datetime(2026, 2, 1, 9), None, datetime(2024, 5, 1, 8), None],
        main.INTAKE_LEGACY: [date(2026, 1, 20), date(2023, 3, 4), date(2024, 4, 1), None],
    })
    out = main.add_intake_date(df)
    assert out[main.INTAKE_DATE].tolist() == [
        datetime(2026, 2, 1, 9),        # stamp present: stamp, even with a legacy date
        date(2023, 3, 4),               # stamp blank: back-filled
        datetime(2024, 5, 1, 8),        # stamp wins over an earlier legacy date
        None]                           # neither: blank
    assert out["intake_date_source"].tolist() == [
        main.INTAKE_ENTERED, main.INTAKE_LEGACY, main.INTAKE_ENTERED, None]

def test_headers_are_english_ascii():
    for header in main.FIXED_HEADERS.values():
        assert header.isascii(), header


def test_workbook_round_trips_every_type():
    import io
    import openpyxl
    import pandas as pd
    df = pd.DataFrame({
        "Text": ["=1+1", None],
        "When": [datetime(2026, 3, 1, 4, 5, 6), None],
        "Day": [date(2026, 3, 5), None],
        "Count": pd.array([7, None], dtype="Int64"),
        "Fee": [5500.5, None],
        "Closed": [True, False],
    })
    ws = openpyxl.load_workbook(io.BytesIO(main.build_workbook(df))).active
    assert ws.title == "Deals"
    assert [c.value for c in ws[1]] == list(df.columns)
    first = [c.value for c in ws[2]]
    assert first[0] == "=1+1" and ws["A2"].data_type == "s"   # text, not a formula
    assert first[1] == datetime(2026, 3, 1, 4, 5, 6)
    assert first[2] == datetime(2026, 3, 5)
    assert first[3] == 7 and first[4] == 5500.5 and first[5] is True
    assert [c.value for c in ws[3]] == [None, None, None, None, None, False]


def test_check_stored_allows_sharepoint_metadata_but_not_stale_or_truncated():
    from datetime import timezone
    started = datetime(2026, 9, 23, 3, 50, tzinfo=timezone.utc)
    fresh = "2026-09-23T03:50:12Z"
    # The real case: SharePoint added ~9 KB of its own metadata.
    assert main.check_stored({"size": 3003318, "lastModifiedDateTime": fresh}, 2994000, started) is None
    assert "stored" in main.check_stored({"size": 1000, "lastModifiedDateTime": fresh}, 2994000, started)
    assert "before" in main.check_stored({"size": 2994000, "lastModifiedDateTime": "2026-09-22T10:00:00Z"},
                                         2994000, started)
    assert main.check_stored({}, 2994000, started) is not None


def test_settled_date_closes_a_deal_in_a_retained_stage():
    stages = [{"id": "13404174", "label": "Retained - Pre Lit", "displayOrder": 8},
              {"id": "5411633", "label": "Intake", "displayOrder": 0}]
    deals = [
        {"id": "1", "properties": {"dealstage": "13404174", "date___settled": "2026-09-15"}},
        {"id": "2", "properties": {"dealstage": "13404174"}},
        {"id": "3", "properties": {"dealstage": "5411633"}},
    ]
    defs = {"date___settled": {"type": "date"}}
    df = main.build_deals_frame(deals, ["dealstage", "date___settled"], defs,
                                {s["id"]: s["label"] for s in stages}, {}, "Lemon Law")
    df = main.add_stage_attributes(df, stages)
    df = main.add_close_date(df, deals, stages)
    assert df["is_settled"].tolist() == [True, False, False]
    assert df["dealstage__closed"].tolist() == [True, False, False]
    assert df["close_date"].tolist() == [date(2026, 9, 15), None, None]
    assert df["close_date_source"].tolist() == ["date___settled", None, None]


def test_fetch_deals_splits_a_window_over_the_search_ceiling():
    """A window whose total reaches 10,000 is halved until every half fits,
    and every deal still comes back exactly once."""
    start = P.localize(datetime(2023, 3, 1))
    end = P.localize(datetime(2023, 4, 1))
    # 12,000 fake deals spread evenly over March.
    span = (end - start).total_seconds()
    created = {str(i): start.timestamp() + span * i / 12000 for i in range(12000)}

    class Resp:
        status_code = 200
        def __init__(self, body): self._b = body
        def json(self): return self._b

    calls = []
    def fake(method, url, headers=None, json=None, **kw):
        f = json["filterGroups"][0]["filters"]
        lo, hi = int(f[1]["value"]) / 1000, int(f[2]["value"]) / 1000
        ids = sorted(i for i, t in created.items() if lo <= t < hi)
        offset = int(json.get("after") or 0)
        page = ids[offset:offset + json["limit"]]
        calls.append((lo, hi))
        nxt = offset + len(page)
        body = {"total": len(ids),
                "results": [{"id": i, "properties": {"createdate": "x", "empty": None}} for i in page]}
        if nxt < len(ids):
            body["paging"] = {"next": {"after": str(nxt)}}
        return Resp(body)

    real = main.request_with_retry
    main.request_with_retry = fake
    try:
        deals = main.fetch_deals(["createdate"], [(start, end)], {})
    finally:
        main.request_with_retry = real
    assert len(deals) == 12000
    assert len({d["id"] for d in deals}) == 12000
    assert all("empty" not in d["properties"] for d in deals)    # nulls dropped
    assert len({c for c in calls}) >= 3                            # split at least once


def test_window_label():
    assert main.window_label(P.localize(datetime(2000, 1, 1)), P.localize(datetime(2021, 1, 1))) \
        == "before 2021-01-01"
    assert main.window_label(P.localize(datetime(2023, 3, 1)), P.localize(datetime(2023, 4, 1))) == "2023-03"
    assert main.window_label(P.localize(datetime(2023, 12, 1)), P.localize(datetime(2024, 1, 1))) == "2023-12"
    assert ".." in main.window_label(P.localize(datetime(2023, 3, 1)), P.localize(datetime(2023, 3, 16)))


def test_unique_headers():
    assert main.unique_headers([("Close Out", "a"), ("Close Out", "b"), ("X", "c")]) == \
        ["Close Out (a)", "Close Out (b)", "X"]


if __name__ == "__main__":
    tests = [v for k, v in dict(globals()).items() if k.startswith("test_")]
    for t in tests:
        t()
        print(f"ok  {t.__name__}")
    print(f"{len(tests)} passed")
