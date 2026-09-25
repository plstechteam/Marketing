"""Tests for splice.py — replacing Deals without touching other tabs.
Run: python tests/test_splice.py"""
import io
import os
import sys
import zipfile
from datetime import date, datetime

import numpy as np
import openpyxl
import pandas as pd
import xlsxwriter

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
import splice  # noqa: E402


def fixture(with_calc_chain=False):
    """Deals + a user tab "Maz" with formulas reading Deals, formatting,
    a chart and shared strings — what the live workbook looks like."""
    b = io.BytesIO()
    wb = xlsxwriter.Workbook(b, {"in_memory": True})
    d = wb.add_worksheet("Deals")
    m = wb.add_worksheet("Maz")
    d.write_row(0, 0, ["Record ID", "Deal Stage", "Fee"])
    for r in range(1, 4):
        d.write_row(r, 0, [f"old{r}", "Close Out", 100 * r])
    m.write("A1", "Rows in Deals")
    m.write_formula("B1", "=COUNTA(Deals!A:A)-1", None, 3)
    m.write("A3", "Nota de Maz, no tocar")
    m.write("A5", "Formato", wb.add_format({"bold": True, "bg_color": "#FFFF00"}))
    chart = wb.add_chart({"type": "column"})
    chart.add_series({"values": "=Deals!$C$2:$C$4"})
    m.insert_chart("D2", chart)
    wb.close()
    data = b.getvalue()
    if not with_calc_chain:
        return data
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for info in src.infolist():
            body = src.read(info.filename)
            if info.filename == "[Content_Types].xml":
                body = body.replace(b"</Types>", b'<Override PartName="/xl/calcChain.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/></Types>')
            if info.filename == "xl/_rels/workbook.xml.rels":
                body = body.replace(b"</Relationships>", b'<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain" Target="calcChain.xml"/></Relationships>')
            z.writestr(info, body)
        z.writestr("xl/calcChain.xml", b'<?xml version="1.0"?><calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><c r="B1" i="2"/></calcChain>')
    return out.getvalue()


NEW = pd.DataFrame({
    "Record ID": ["n1", "n2", "n3", "n4", "n5"],
    "Deal Stage": ["Settled", " lead space", "=HYPERLINK(1)", "x&<y>", None],
    "Fee": [5500.5, np.float64(10), None, 7, 1],
    "When": [datetime(2026, 3, 1, 4, 5, 6), None, None, None, None],
    "Day": [date(2026, 3, 5)] + [None] * 4,
    "Days": pd.array([1, None, 3, 4, 5], dtype="Int64"),
    "Closed": np.array([True, False, True, False, True]),
})


def test_only_deals_changes_and_maz_is_byte_identical():
    before = fixture()
    after, rep = splice.replace_sheet(before, "Deals", NEW)
    assert rep["sheets"] == ["Deals", "Maz"]
    assert rep["changed"] == ["xl/styles.xml", "xl/workbook.xml", rep["sheet_part"]]
    assert splice.untouched_parts_identical(before, after, rep)
    for part in ("xl/worksheets/sheet2.xml", "xl/charts/chart1.xml", "xl/drawings/drawing1.xml",
                 "xl/sharedStrings.xml", "xl/worksheets/_rels/sheet2.xml.rels"):
        assert part in rep["untouched"], part


def test_new_deals_data_reads_back_with_the_right_types():
    after, _ = splice.replace_sheet(fixture(), "Deals", NEW)
    wb = openpyxl.load_workbook(io.BytesIO(after))
    deals, maz = wb["Deals"], wb["Maz"]
    assert [c.value for c in deals[1]] == list(NEW.columns)
    assert [c.value for c in deals[2]] == [
        "n1", "Settled", 5500.5, datetime(2026, 3, 1, 4, 5, 6), datetime(2026, 3, 5), 1, True]
    assert deals["B3"].value == "lead space"      # trimmed
    assert deals["B4"].value == "=HYPERLINK(1)" and deals["B4"].data_type == "s"  # text, not formula
    assert deals["B5"].value == "x&<y>"
    assert deals.max_row == 6                      # old 4 rows gone, 5 + header
    assert maz["B1"].value == "=COUNTA(Deals!A:A)-1"
    assert maz["A3"].value == "Nota de Maz, no tocar"
    assert maz["A5"].font.b is True


def test_rerun_is_stable_and_recalc_is_forced():
    once, _ = splice.replace_sheet(fixture(), "Deals", NEW)
    twice, _ = splice.replace_sheet(once, "Deals", NEW)
    read = lambda b, n: zipfile.ZipFile(io.BytesIO(b)).read(n)  # noqa: E731
    assert read(once, "xl/styles.xml") == read(twice, "xl/styles.xml")
    assert b'fullCalcOnLoad="1"' in read(once, "xl/workbook.xml")
    assert read(twice, "xl/workbook.xml").count(b"fullCalcOnLoad") == 1


def test_calc_chain_is_removed_with_its_references():
    after, rep = splice.replace_sheet(fixture(with_calc_chain=True), "Deals", NEW)
    z = zipfile.ZipFile(io.BytesIO(after))
    assert "xl/calcChain.xml" not in z.namelist()
    assert b"calcChain" not in z.read("[Content_Types].xml")
    assert b"calcChain" not in z.read("xl/_rels/workbook.xml.rels")
    assert rep["dropped"] == ["xl/calcChain.xml"]
    openpyxl.load_workbook(io.BytesIO(after))      # still opens


def test_refuses_when_deals_is_missing_or_has_attached_objects():
    for name in ("Nope",):
        try:
            splice.replace_sheet(fixture(), name, NEW)
        except splice.SpliceError:
            pass
        else:
            raise AssertionError("a missing sheet must refuse")
    try:
        splice.replace_sheet(fixture(), "Maz", NEW)   # Maz has a chart attached
    except splice.SpliceError as exc:
        assert "attached" in str(exc)
    else:
        raise AssertionError("a sheet with attached objects must refuse")


def test_shift_guard_catches_a_column_inserted_in_the_middle_only():
    before = fixture()           # Maz: formula on Deals!A:A, chart on Deals!$C$2:$C$4
    refs = splice.referenced_columns(before, "Deals")
    assert set(refs) == {"A", "C"}
    inserted = pd.DataFrame({"Record ID": ["a"], "New": ["n"], "Deal Stage": ["x"], "Fee": [1]})
    after, _ = splice.replace_sheet(before, "Deals", inserted)
    assert [s[0] for s in splice.shifted_references(before, after, "Deals")] == ["C"]
    appended = pd.DataFrame({"Record ID": ["a"], "Deal Stage": ["x"], "Fee": [1], "New": ["n"]})
    after, _ = splice.replace_sheet(before, "Deals", appended)
    assert splice.shifted_references(before, after, "Deals") == []
    assert splice.header_row(after, "Deals") == {"A": "Record ID", "B": "Deal Stage", "C": "Fee", "D": "New"}


def test_column_letters():
    assert [splice.column_letter(i) for i in (0, 25, 26, 51, 52, 701, 702)] == \
        ["A", "Z", "AA", "AZ", "BA", "ZZ", "AAA"]


def test_blank_text_is_an_empty_cell():
    df = pd.DataFrame({"Record ID": ["a", "b", "c", "d", "e"],
                       "Note": ["", "   ", "\u00a0\u200b", " x\u00a0", None],
                       "Amount": [None, np.nan, 0.0, 12.5, None]})
    after, _ = splice.replace_sheet(fixture(), "Deals", df)
    xml = zipfile.ZipFile(io.BytesIO(after)).read("xl/worksheets/sheet1.xml").decode()
    for ref in ("B2", "B3", "B4", "B6", "C2", "C3", "C6"):
        assert f'r="{ref}"' not in xml, ref     # no cell at all, not an empty string
    ws = openpyxl.load_workbook(io.BytesIO(after))["Deals"]
    assert ws["B5"].value == "x" and ws["C4"].value == 0 and ws["C5"].value == 12.5


if __name__ == "__main__":
    tests = [v for k, v in dict(globals()).items() if k.startswith("test_")]
    for t in tests:
        t()
        print(f"ok  {t.__name__}")
    print(f"{len(tests)} passed")
